import os
import openpyxl
from pathlib import Path
import pandas as pd
import warnings


# write the absolute location of filetanks to the path specified
def write_tanks_to_excel(tank_path, excel_path):
    all_paths = []

    head_tail = os.path.split(excel_path)
    check_folder = os.path.isdir(head_tail[0])
    if head_tail[0] is not '' and not check_folder:
        os.makedirs(head_tail[0])

    for root, dirs, files in os.walk(tank_path):
        # todo: check length of folder structures
        if len(files) > 2 and (files[2].endswith('tev') or files[3].endswith('tev')):
            all_paths.append(root)

    final_paths = [i for i in all_paths if "Bad" not in i and "bad" not in i]

    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active

    for idx, val in enumerate(final_paths):
        ws.cell(row=idx + 2, column=1, value=val)

    wb.save(excel_path)


def get_list(file_path, target_column):  # in our case, condition=4, stim=4
    d = []
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active
    # num_recordings = ws.max_row - ws.min_row
    num_recordings = 0
    cols = tuple(ws.columns)
    max_length = len(cols[0])

    # Check last recording in custom column. Assuming list of target directories is shorter than length of excel file
    for row in range(1, max_length + 1):
        if ws.cell(row=row, column=target_column).value is None:
            num_recordings = row - 2
            break
    if ws.cell(row=max_length, column=target_column).value is not None:
        num_recordings = max_length - 1

    for i in range(num_recordings):
        d.append(ws.cell(i + 2, target_column).value)

    return d


def create_dictionary(file_path):
    """
    Read file or folder names from column A and values from columns B-D.

    Searches the experiment directory and all subdirectories for matching
    files or folders. The returned dictionary uses the matched Path objects
    as keys.

    Expected workbook location:
        experiment_directory / Notes / workbook.xlsx

    Returns:
        {
            Path("path/to/file_or_folder"): [column_b, column_c, column_d],
            ...
        }
    """
    file_path = Path(file_path)

    if not file_path.is_file():
        raise FileNotFoundError(
            f"Excel file does not exist: {file_path}"
        )

    # Move from:
    # experiment_directory / Notes / workbook.xlsx
    # to:
    # experiment_directory
    data_folder = file_path.parents[1]

    workbook = openpyxl.load_workbook(
        file_path,
        data_only=True
    )
    worksheet = workbook.active

    # Index both files and directories by name.
    # Using casefold() makes matching case-insensitive.
    path_index = {}

    for path in data_folder.rglob("*"):
        path_index.setdefault(
            path.name.casefold(),
            []
        ).append(path)

    path_dictionary = {}

    for row in worksheet.iter_rows(
        min_row=2,
        min_col=1,
        max_col=4,
        values_only=True
    ):
        entry_name, id_1, id_2, id_3 = row

        # Stop at the first blank cell in column A
        if entry_name is None:
            break

        entry_name = str(entry_name).strip()

        matches = path_index.get(
            entry_name.casefold(),
            []
        )

        if not matches:
            warnings.warn(
                f"File or folder not found: {entry_name}",
                category=UserWarning,
                stacklevel=2
            )
            continue

        if len(matches) > 1:
            warnings.warn(
                f"Multiple files or folders named {entry_name!r} "
                f"were found. All matches will be included.",
                category=UserWarning,
                stacklevel=2
            )

        for matched_path in matches:
            path_dictionary[matched_path] = [
                id_1,
                id_2,
                id_3
            ]

    return path_dictionary

def find_files(file_dictionary, data_folder):
    """
    Recursively search data_folder for filenames contained in file_dictionary.

    Missing files generate warnings, but do not stop execution.

    Returns:
        Dictionary mapping each class name to the files that were found.
    """
    data_folder = Path(data_folder)

    if not data_folder.is_dir():
        raise NotADirectoryError(
            f"The data folder does not exist: {data_folder}"
        )

    # Index every file in the data folder and its subfolders
    file_index = {}

    for path in data_folder.rglob("*"):
        if path.is_file():
            file_index.setdefault(path.name.lower(), []).append(path)

    found_files = {}

    for class_name, filenames in file_dictionary.items():
        found_files[class_name] = {}

        for filename in filenames:
            # Ignore blank Excel cells
            if filename is None:
                continue

            filename = str(filename).strip()
            matches = file_index.get(filename.lower(), [])

            if matches:
                found_files[class_name][filename] = matches
            else:
                warnings.warn(
                    f"File not found for {class_name!r}: {filename}",
                    category=UserWarning,
                    stacklevel=2
                )

    return found_files


def create_legends(file_path, **kwargs):
    type_nums = []
    for key, value in kwargs.items():
        type_nums.append(value)

    wb = openpyxl.load_workbook(file_path)
    ws = wb.active
    legend = []
    iterator = 0
    for my_type in type_nums:
        row_names = []
        for i in range(my_type):
            row_names.append(ws.cell(2 + iterator, 2).value)
            iterator += 1
        legend.append(row_names)
    return legend

def resolve_ephys_paths_from_excel(excel_path):
    """
    Resolve ephys data paths using Excel file location.

    Assumes structure:
    /<date>/Notes/<excel>
    /<date>/Data/<experiment>/<recording_folder>
    """

    excel_path = Path(excel_path)

    # --- derive base_dir relative to excel ---
    # Notes → parent (date folder) → Data
    base_dir = excel_path.parent.parent / "Data"

    if not base_dir.exists():
        raise FileNotFoundError(f"Derived Data directory not found: {base_dir}")

    print(f"Using base_dir: {base_dir}")

    # --- load excel ---
    df = pd.read_excel(excel_path)

    if "Location" not in df.columns:
        raise ValueError("Excel must contain a 'Location' column")

    # --- build directory lookup ---
    print("Scanning directory tree (one-time cost)...")
    all_dirs = {}
    duplicates = {}

    for p in base_dir.rglob("*"):
        if p.is_dir():
            name = p.name
            if name in all_dirs:
                duplicates.setdefault(name, []).append(p)
            else:
                all_dirs[name] = p

    if duplicates:
        print(f"⚠️ Found duplicate folder names ({len(duplicates)}). Using first match.")

    # --- resolve paths ---
    resolved_paths = []
    missing = []

    for loc in df["Location"]:
        p = all_dirs.get(loc)

        if p is None:
            missing.append(loc)
            resolved_paths.append(None)
        else:
            resolved_paths.append(p)

    df["full_path"] = resolved_paths

    # --- reporting ---
    print(f"\nResolved: {df['full_path'].notna().sum()} / {len(df)}")

    if missing:
        print(f"\n⚠️ Missing {len(missing)} folders (showing up to 10):")
        for m in missing[:10]:
            print("   ", m)

    return df, [p for p in resolved_paths if p is not None]

# # create_dictionary("Surgical Log.xlsx")
# tdt_block_path = "C:\\Users\\steph\\PycharmProjects\\pyeCAP\\pyeCAP\\TestData\\20200113_SLB_ImtheraVNS\\TDT"
# excel_path = ("new Directory2/Surgical Log.xlsx")
#
# write_tanks_to_excel(tdt_block_path, excel_path)
# create_dictionary(excel_path)
# example_surgical_log_path="D:\\20191216_SLB_ImtheraVNS\\Surgical Log.xlsx"
# a = create_labels(example_surgical_log_path, mono=6, bi=4)
